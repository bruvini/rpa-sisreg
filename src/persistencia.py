"""Atualização transacional da XLSX original com backup, hash e escrita atômica.

O backup é mantido no mesmo diretório institucional da origem. Nenhum arquivo
com dados de saúde é enviado a serviços externos por esta rotina.
"""
from __future__ import annotations
import hashlib
import os
import shutil
import tempfile
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook

COLUNAS_AUDITORIA = ('rpaSituacao', 'rpaDetalhe', 'rpaOrigemPDF',
                   'rpaConsultaAmbulatorial', 'rpaAtualizadoEm',
                   'rpaConsultaHospitalar', 'rpaProcedimentoSISREG',
                   'rpaTempoConsultaSegundos', 'rpaNumeroSolicitacao')

class AlteracaoExterna(Exception):
    pass


def sha256(caminho: Path) -> str:
    h = hashlib.sha256()
    with caminho.open('rb') as f:
        for bloco in iter(lambda: f.read(1024 * 1024), b''):
            h.update(bloco)
    return h.hexdigest()


def gravar_na_origem(sessao, *, motivo='preparacao', backup=True) -> dict:
    """Salva somente campos vazios e trilha; nunca substitui conflitos.

    Detecta edições feitas fora do RPA desde a última leitura/salvamento;
    recria o arquivo via substituição atômica em mesmo diretório.
    """
    if sessao.origem is None:
        raise ValueError('Selecione uma planilha XLSX antes de gravar.')
    origem = Path(sessao.origem)
    if not origem.exists():
        raise FileNotFoundError('A planilha hospitalar não está mais disponível.')
    atual = sha256(origem)
    if atual != sessao.hash_planilha:
        raise AlteracaoExterna('Planilha alterada externamente. Recarregue antes de salvar; nada foi sobrescrito.')

    wb = load_workbook(origem, keep_links=False, data_only=False)
    quantidade_cpf = quantidade_cns = quantidade_status = quantidade_hospital = 0
    try:
        for aba in ('POA CIRURGICO','PT6','ALTA COMPLEXIDADE'):
            ws = wb[aba]
            for col,cab in enumerate(COLUNAS_AUDITORIA, 8):
                if ws.cell(1,col).value not in (None, cab):
                    raise AlteracaoExterna(f'Coluna de auditoria ocupada em {aba}, {col}. Não sobrescrevemos.')
                ws.cell(1,col).value = cab
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = f'A1:P{max(2, ws.max_row)}'
        for r in sessao.registros:
            ws = wb[r.aba]
            cpf = ws.cell(r.linha,4)
            cns = ws.cell(r.linha,5)
            # Não alterar entradas preexistentes, nem aplicar valores não verificados.
            if not r.cpf_origem and r.cpf and cpf.value in ('', None):
                cpf.value = str(r.cpf)
                cpf.number_format = '@'
                quantidade_cpf += 1
            if not r.cns_origem and r.cns and cns.value in ('',None):
                cns.value = str(r.cns)
                cns.number_format = '@'
                quantidade_cns += 1
            carater=ws.cell(r.linha,6)
            autorizador=ws.cell(r.linha,7)
            if r.carater and carater.value in ('',None):
                carater.value=str(r.carater);carater.number_format='@'
                quantidade_hospital+=1
            if r.autorizador and autorizador.value in ('',None):
                autorizador.value=str(r.autorizador);autorizador.number_format='@'
                quantidade_hospital+=1
            valores=(r.situacao, r.detalhe, r.origem_pdf,
                     r.consulta_ambulatorial, r.atualizado_em, r.consulta_hospitalar,
                     r.procedimento_sisreg, str(r.tempo_consulta) if r.tempo_consulta else '',
                     r.solicitacao)
            for col,valor in enumerate(valores,8):
                valor = str(valor or '')
                if valor.startswith(('=','+','-','@')):
                    valor = "'" + valor
                cel=ws.cell(r.linha,col)
                # Não apagar solicitação de uma execução anterior sem nova evidência.
                if col == 16 and not valor and cel.value not in ('', None):
                    continue
                if cel.value != valor:
                    cel.value = valor
                    quantidade_status += 1
        # Cria arquivo temporário dentro do mesmo volume para os.replace ser atômico.
        fd,nome_tmp=tempfile.mkstemp(prefix='.rpa_sisreg_',suffix='.xlsx',dir=origem.parent)
        os.close(fd)
        temporario=Path(nome_tmp)
        try:
            wb.save(temporario)
            # Validar integridade do ZIP/XLSX antes de substituir a origem.
            verificar=load_workbook(temporario,read_only=True,data_only=False,keep_links=False)
            try:
                for aba in ('POA CIRURGICO','PT6','ALTA COMPLEXIDADE'):
                    if aba not in verificar.sheetnames:
                        raise ValueError('Falta aba no arquivo atualizado: '+aba)
            finally:
                verificar.close()
            if sha256(origem) != sessao.hash_planilha:
                raise AlteracaoExterna('A planilha foi modificada enquanto o RPA a preparava.')
            copia_backup = origem.with_name(origem.stem + '.backup_antes_RPA.xlsx')
            if backup and not getattr(sessao, "fonte_original", None) and not copia_backup.exists():
                # A primeira versão do arquivo permanece recuperável.
                shutil.copy2(origem,copia_backup)
            os.replace(temporario,origem)
            sessao.hash_planilha = sha256(origem)
        finally:
            temporario.unlink(missing_ok=True)
    finally:
        wb.close()
    return {'cpf':quantidade_cpf,'cns':quantidade_cns,'hospital':quantidade_hospital,
            'campos_atualizados':quantidade_status,
            'backup':str(origem.with_name(origem.stem+'.backup_antes_RPA.xlsx')),
            'motivo':motivo}
