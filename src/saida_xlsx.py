"""Apresentação XLSX da fiscalização — não modifica conclusões do motor.

Estados técnicos ficam na aplicação e no diagnóstico sanitizado; na planilha,
resultados apresentam a situação da trilha PDF → cadastro → internação.
"""
from __future__ import annotations
from datetime import datetime
from aih_extract import digitos

CAMPOS_FINAIS = ('rpaSituacao', 'rpaAtualizadoEm', 'rpaProcedimentoSISREG', 'rpaNumeroSolicitacao')
CABECALHOS_LEGADOS = ('rpaSituacao', 'rpaDetalhe', 'rpaOrigemPDF',
    'rpaConsultaAmbulatorial', 'rpaAtualizadoEm', 'rpaConsultaHospitalar',
    'rpaProcedimentoSISREG', 'rpaTempoConsultaSegundos', 'rpaNumeroSolicitacao')

AUSENCIA_PDF = {'AIH_AUSENTE_NO_PDF', 'ESPELHO_NAO_LOCALIZADO', 'AIH_FORMATO_INVALIDO'}
DIVERGENCIA_PDF = {'NOME_DIVERGENTE_ESPELHO', 'NOME_AUSENTE_NO_ESPELHO',
    'ESPELHOS_CONFLITANTES', 'CPF_CONFLITANTE_PDF', 'CNS_CONFLITANTE_PDF',
    'CPF_HOSPITAL_INVALIDO', 'CNS_HOSPITAL_INVALIDO', 'NOME_NAO_INFORMADO',
    'DOCUMENTO_NAO_CLASSIFICADO', 'PDF_SEM_DOCUMENTO_VALIDO'}


def situacao_final(registro) -> str:
    """Prioriza o primeiro impedimento material, sem inventar consulta concluída.

    Uma busca que falhou tecnicamente não deve parecer ausência comprovada de AIH.
    """
    situacao = str(registro.situacao or '').upper()
    if situacao in AUSENCIA_PDF or (not registro.cpf and situacao.startswith('AIH_')):
        return 'Não encontrado no espelho de AIH'
    if situacao in DIVERGENCIA_PDF:
        return 'Dados divergentes do espelho de AIH'
    if situacao == 'DADOS_COLETADOS_SEM_DIVERGENCIA':
        return 'Consulta sem divergências'
    if situacao.startswith('PENDENTE_'):
        motivo = situacao.removeprefix('PENDENTE_')
        if any(x in motivo for x in ('PROCEDIMENTO_DIVERGENTE', 'CARATER_CONFLITANTE', 'AUTORIZADOR_CONFLITANTE')):
            return 'Dados divergentes no SISREG'
        if any(x in motivo for x in ('AIH_NAO_COMPROVADA', 'FICHA_NAO_ABERTA')):
            return 'Não encontrado AIH no SISREG' if motivo.startswith('AIH_') else 'Consulta não concluída — revisão necessária'
        if any(x in motivo for x in ('CPF_DIVERGENTE', 'NOME_DIVERGENTE', 'CNS_CONFLITO', 'CNS_DIVERGENTE')):
            return 'Dados divergentes na consulta SISREG'
        if any(x in motivo for x in ('NOME_NAO_CONFIRMADO', 'CNS_NAO_IDENTIFICADO', 'CNS_NAO_UNICO',
                                    'FICHA_CADASTRO_NAO_IDENTIFICADA', 'CPF_NAO_CONFIRMADO')):
            return 'Não encontrado na consulta SISREG'
        if any(x in motivo for x in ('FALHA_', 'CAMPOS_FICHA_', 'TIMEOUT', 'NAVEGACAO')):
            return 'Consulta não concluída — revisão necessária'
        return 'Consulta não concluída — revisão necessária'
    if 'CONFLITO' in situacao or 'DIVERGENTE' in situacao:
        return 'Dados divergentes na consulta SISREG'
    if situacao in ('CNS_CONFIRMADO_SISREG', 'CNS_CONFIRMADO_OPERADORA',
                    'DOCUMENTO_PDF_LOCALIZADO', 'DOCUMENTO_VALIDADO_POR_AIH_E_NOME',
                    'PRONTO_PARA_CONFERENCIA', 'AGUARDA_PREPARACAO'):
        return 'Aguardando consulta hospitalar' if registro.cns else 'Aguardando consulta SISREG'
    return 'Consulta não concluída — revisão necessária'


def comparacao_procedimento(registro) -> str:
    """Compara códigos, jamais infere conformidade por ausência de um dos lados."""
    hospital = digitos(str(registro.procedimento or ''))
    sisreg = digitos(str(registro.procedimento_sisreg or ''))
    if len(hospital) != 10 or len(sisreg) != 10:
        return 'Não conferido'
    return 'Compatível' if hospital == sisreg else 'Divergente'


def data_br(valor: str):
    if not valor: return None
    try:
        return datetime.fromisoformat(str(valor))
    except ValueError:
        return None


def adequar_colunas(ws) -> None:
    """Preserva A:G e outras colunas não controladas; migra só layout conhecido."""
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    cabecalhos = [str(ws.cell(1,c).value or '') for c in range(8, max(ws.max_column, 16)+1)]
    existentes = [c for c in cabecalhos if c]
    # Outras colunas não são nossas — não deletar dados do hospital.
    if existentes and tuple(existentes[:len(CAMPOS_FINAIS)]) != CAMPOS_FINAIS:
        if tuple(existentes[:len(CABECALHOS_LEGADOS)]) == CABECALHOS_LEGADOS:
            ws.delete_cols(8, len(CABECALHOS_LEGADOS))
        else:
            raise ValueError('Colunas adicionais não reconhecidas: manter dados e revisar a planilha antes de exportar.')
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    for col, titulo in enumerate(CAMPOS_FINAIS, 8):
        c=ws.cell(1,col)
        c.value=titulo
        c.fill=PatternFill('solid',fgColor='083F6C')
        c.font=Font(name='Aptos',size=10,bold=True,color='FFFFFF')
        c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        c.border=Border(bottom=Side(style='medium',color='077BA1'))
    # Destacar a origem da informação para impedir interpretação como decisão fiscal.
    from openpyxl.comments import Comment
    ws.cell(1,8).comment=Comment('Classificação automatizada dos campos disponíveis. NÃO constitui atesto fiscal.','RPA SISREG')
    ws.cell(1,10).comment=Comment('Comparação entre código hospitalar e código apresentado na ficha SISREG.','RPA SISREG')
    ws.cell(1,11).comment=Comment('Número da SOLICITAÇÃO extraído da ficha SISREG mediante correspondência da AIH.','RPA SISREG')
    ws.cell(1,7).comment=Comment('Código autorizador e código de solicitação não são equivalentes por definição. Quando G estiver inicialmente vazia, recebe o Código Solicitação de ficha SISREG com AIH confirmada, por regra operacional solicitada pela gestão. Não constitui comprovação de um autorizador independente. Número da solicitação também é informado em K.','RPA SISREG')
    larguras={'A':32,'B':20,'C':21,'D':18,'E':19,'F':19,'G':24,'H':43,'I':23,'J':22,'K':25}
    for col,largura in larguras.items():ws.column_dimensions[col].width=largura
    ws.row_dimensions[1].height=34
    ws.freeze_panes='A2'
    ws.auto_filter.ref=f'A1:K{max(2,ws.max_row)}'
    # Modelos oficiais contêm Tabelas do Excel (A:G). Estender a tabela real,
    # em vez de somente as células, mantém filtros e faixas visuais coerentes.
    from openpyxl.worksheet.table import TableColumn
    from openpyxl.worksheet.filters import AutoFilter
    for tabela in ws.tables.values():
        if not tabela.ref.startswith('A1:'):
            continue
        nova_faixa=f'A1:K{max(2,ws.max_row)}'
        tabela.ref=nova_faixa
        cabecalhos=[str(ws.cell(1,col).value) for col in range(1,12)]
        tabela.tableColumns=[TableColumn(id=i,name=nome) for i,nome in enumerate(cabecalhos,1)]
        tabela.autoFilter=AutoFilter(ref=nova_faixa)
    ws.sheet_view.showGridLines=False


def escrever_linha(ws, registro) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    estado=situacao_final(registro)
    data=data_br(registro.atualizado_em)
    comparacao=comparacao_procedimento(registro)
    valores=(estado,data,comparacao,registro.solicitacao or None)
    color={'Consulta sem divergências':'13704A', 'Dados divergentes no SISREG':'9C4319'}
    for coluna,valor in enumerate(valores,8):
        cel=ws.cell(registro.linha,coluna)
        cel.value=valor
        cel.font=Font(name='Aptos',size=10,color=color.get(estado,'293D52') if coluna==8 else '293D52',bold=(coluna==8))
        cel.alignment=Alignment(vertical='center',wrap_text=coluna==8)
        cel.fill=PatternFill('solid',fgColor='F1F7FA' if registro.linha%2==0 else 'FFFFFF')
    ws.cell(registro.linha,9).number_format='DD/MM/YYYY HH:MM:SS'
    ws.cell(registro.linha,11).number_format='@'
    ws.row_dimensions[registro.linha].height=28
