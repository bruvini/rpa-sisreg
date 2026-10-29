"""Regressão da preparação v0.7.3, apenas com planilhas sintéticas."""
import os
import sys
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from planilha import SessaoPlanilha
from saida_xlsx import CAMPOS_FINAIS, CABECALHOS_LEGADOS, LayoutPlanilhaIncompativel
from estatisticas import categoria_excecao
from app_info import APP_VERSION
from openpyxl import load_workbook

FIXTURE=Path(__file__).parent / 'fixtures' / 'hospital_sintetico.xlsx'

class TesteMigracaoLegada(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.fonte=Path(self.temp.name)/'origem.xlsx'
        shutil.copy2(FIXTURE,self.fonte)

    def construir_legado(self,quantidade):
        wb=load_workbook(self.fonte)
        for aba in ('POA CIRURGICO','PT6','ALTA COMPLEXIDADE'):
            ws=wb[aba]
            for col,header in enumerate(CABECALHOS_LEGADOS[:quantidade],8):
                ws.cell(1,col,header)
            ws['H2']='DOCUMENTO_PDF_LOCALIZADO'
            ws['I2']='Diagnóstico fictício'
        wb.save(self.fonte)
        wb.close()

    def validar_resultado(self):
        s=SessaoPlanilha();s.carregar(self.fonte)
        s.validar_layout_para_saida()
        antes=self.fonte.read_bytes()
        copia=s.preparar_copia_de_trabalho()
        gravado=s.salvar_na_planilha('teste_sintetico')
        self.assertEqual(self.fonte.read_bytes(),antes)
        self.assertTrue(copia.is_file())
        self.assertEqual(gravado['campos_atualizados'],len(s.registros))
        wb=load_workbook(copia,read_only=True)
        try:
            for aba in ('POA CIRURGICO','PT6','ALTA COMPLEXIDADE'):
                ws=wb[aba]
                self.assertEqual(tuple(ws.cell(1,i).value for i in range(8,11)),CAMPOS_FINAIS)
                self.assertIsNone(ws.cell(1,12).value)
                for reg in (r for r in s.registros if r.aba == aba):
                    self.assertIsNotNone(ws.cell(reg.linha,8).value)
        finally:wb.close()

    def test_variante_de_oito_colunas_que_causava_falha(self):
        self.construir_legado(8)
        self.validar_resultado()

    def test_variante_de_nove_colunas(self):
        self.construir_legado(9)
        self.validar_resultado()

    def test_legado_oito_colunas_nao_apaga_dado_sem_cabecalho_em_p(self):
        self.construir_legado(8)
        wb=load_workbook(self.fonte)
        wb['POA CIRURGICO']['P2']='DadoSinteticoExternoNaoMapeado'
        wb.save(self.fonte);wb.close()
        s=SessaoPlanilha();s.carregar(self.fonte)
        with self.assertRaises(LayoutPlanilhaIncompativel):
            s.validar_layout_para_saida()
        self.assertEqual(len(list(Path(self.temp.name).glob('Planilha_RPA_*.xlsx'))),0)

    def test_coluna_desconhecida_bloqueada_sem_gerar_copia(self):
        wb=load_workbook(self.fonte)
        wb['POA CIRURGICO']['H1']='campoHospitalDesconhecido'
        wb.save(self.fonte);wb.close()
        s=SessaoPlanilha();s.carregar(self.fonte)
        with self.assertRaises(LayoutPlanilhaIncompativel):s.validar_layout_para_saida()
        self.assertEqual(len(list(Path(self.temp.name).glob('Planilha_RPA_*.xlsx'))),0)
        self.assertEqual(categoria_excecao('LayoutPlanilhaIncompativel'),'LAYOUT_XLSX_INCOMPATIVEL')

    def test_sem_modificacao_descarta_copia_incompleta(self):
        s=SessaoPlanilha();s.carregar(self.fonte)
        antes=self.fonte.read_bytes()
        copia=s.preparar_copia_de_trabalho()
        self.assertTrue(copia.exists())
        self.assertTrue(s.descartar_copia_nao_gravada())
        self.assertFalse(copia.exists())
        self.assertEqual(self.fonte.read_bytes(),antes)

    def test_copia_modificada_nunca_sera_descartada(self):
        s=SessaoPlanilha();s.carregar(self.fonte)
        copia=s.preparar_copia_de_trabalho()
        s.salvar_na_planilha('teste')
        self.assertFalse(s.descartar_copia_nao_gravada())
        self.assertTrue(copia.exists())

    def test_versao_eh_consistente(self):
        self.assertEqual(APP_VERSION,'0.7.10')

@unittest.skipUnless(os.environ.get('SISREG_GUI_TEST')=='1','Exige ambiente gráfico')
class TesteLogFalha(unittest.TestCase):
    def test_txt_criado_em_falha_no_preparo_sem_pii(self):
        from sisreg_login import Aplicativo
        import tkinter as tk
        with tempfile.TemporaryDirectory() as diretorio:
            fonte=Path(diretorio)/'dados_sinteticos.xlsx'
            shutil.copy2(FIXTURE,fonte)
            root=tk.Tk()
            try:
                app=Aplicativo(root)
                app.arquivo_planilha=fonte
                app.inicio_ciclo=__import__('time').perf_counter()
                app.estatisticas.concluir('preparo',2,consultados=1,encontrados=0,
                                          pendencias=1,falhas=1,completa=False)
                app.estatisticas.diagnosticar('SALVAMENTO_XLSX','ERRO',segundos=2,
                                             categoria='LAYOUT_XLSX_INCOMPATIVEL')
                with patch('sisreg_login.pastas_de_saida', return_value=(Path(diretorio)/'Análises', Path(diretorio))):
                    app._registrar_falha_preparacao(2)
                arquivos=list(Path(diretorio).glob('RELATORIO_TECNICO_RPA_FALHA_*.txt'))
                self.assertEqual(len(arquivos),1)
                conteudo=arquivos[0].read_text(encoding='utf8')
                self.assertIn('LAYOUT_XLSX_INCOMPATIVEL',conteudo)
                self.assertIn('Situação: PARCIAL',conteudo)
                self.assertNotIn('11144477735',conteudo)
                self.assertNotIn('ANA MARIA',conteudo)
                self.assertNotIn('senha-ficticia',conteudo.lower())
            finally:root.destroy()

if __name__=='__main__':unittest.main()
