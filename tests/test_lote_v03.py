import tempfile
import unittest
import sys
from pathlib import Path
from shutil import copy2
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]/'src'))
from planilha import SessaoPlanilha
from persistencia import AlteracaoExterna
from aih_extract import extrair_texto_paginas
from sisreg_navigation import avaliar_resultado_cadastro

class PlanilhaNaOrigem(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.xlsx=Path(self.tmp.name)/'hospital.xlsx'
        copy2(Path(__file__).parent/'fixtures'/'hospital_sintetico.xlsx',self.xlsx)
    def montar_sessao(self):
        s=SessaoPlanilha();s.carregar(self.xlsx)
        pdf=Path(self.tmp.name)/'simulado.pdf';pdf.write_text('dados puramente sintéticos')
        ix=extrair_texto_paginas(['Num AIH: 900000000012-3\nPaciente: ANA MARIA\nCNS/CPF: 11144477735'])
        # Fixture real may have nome ANA MARIA EXEMPLO; use matching name dynamically.
        r=s.registros[0]
        ix=extrair_texto_paginas([f'Num AIH: {r.aih[:12]}-{r.aih[-1]}\nPaciente: {r.nome}\nCNS/CPF: 11144477735'])
        s.complementar_pdf(ix,pdf)
        return s
    def test_salva_pdf_na_origem_e_backup(self):
        s=self.montar_sessao();r=s.registros[0]
        resultado=s.salvar_na_planilha(motivo='pdf')
        wb=load_workbook(self.xlsx);ws=wb[r.aba]
        self.assertEqual(ws.cell(r.linha,4).value,'11144477735')
        self.assertEqual(ws.cell(1,4).value,'numeroCPF')
        self.assertEqual(len(wb.sheetnames),3)
        wb.close()
        backup=Path(resultado['backup'])
        self.assertTrue(backup.exists())
        antigo=load_workbook(backup);self.assertIsNone(antigo[r.aba].cell(r.linha,4).value);antigo.close()
    def test_grava_cns_e_conserva_backup_original(self):
        s=self.montar_sessao();s.salvar_na_planilha('pdf')
        r=s.registros[0]
        s.registrar_cadastro_em_lote(r,'123456789012345','CNS_CONFIRMADO_TELA')
        s.salvar_na_planilha('lote')
        w=load_workbook(self.xlsx);self.assertEqual(w[r.aba].cell(r.linha,5).value,'123456789012345');w.close()
        b=load_workbook(self.xlsx.with_name(self.xlsx.stem+'.backup_antes_RPA.xlsx'))
        self.assertIsNone(b[r.aba].cell(r.linha,5).value);b.close()
    def test_detecta_alteracao_externa(self):
        s=self.montar_sessao();w=load_workbook(self.xlsx)
        w[s.registros[0].aba].cell(s.registros[0].linha,7).value='ALTERADO_POR_TERCEIRO';w.save(self.xlsx);w.close()
        with self.assertRaises(AlteracaoExterna):s.salvar_na_planilha('pdf')
    def test_nao_substitui_cns_conflitante(self):
        s=self.montar_sessao();r=s.registros[0]
        r.cns='123456789012345'
        s.registrar_cadastro_em_lote(r,'987654321098765','CNS_CONFIRMADO_TELA')
        self.assertEqual(r.cns,'123456789012345')
        self.assertEqual(r.situacao,'CNS_CONFLITO_SISREG')

class CadastroConservador(unittest.TestCase):
    def test_confirmacao_valida(self):
        texto='Nome: ANA MARIA EXEMPLO\nCPF: 111.444.777-35\nCNS: 123456789012345'
        self.assertEqual(avaliar_resultado_cadastro(texto,'Ana Maria Exemplo','11144477735'),('123456789012345','CNS_CONFIRMADO_TELA'))
    def test_nega_nome_divergente(self):
        c,m=avaliar_resultado_cadastro('Nome: OUTRA PESSOA\nCNS: 123456789012345','ANA MARIA','11144477735')
        self.assertIsNone(c)
    def test_nega_cpf_divergente(self):
        c,m=avaliar_resultado_cadastro('Nome: ANA MARIA\nCPF: 999.888.777-66\nCNS: 123456789012345','ANA MARIA','11144477735')
        self.assertEqual(m,'CPF_DIVERGENTE_CADASTRO')
    def test_nega_cns_ambiguos(self):
        c,m=avaliar_resultado_cadastro('Nome: ANA MARIA\nCNS: 123456789012345\nCNS: 987654321098765','ANA MARIA','11144477735')
        self.assertIsNone(c)
    def test_sem_rotulo_cns_nao_autoriza(self):
        c,m=avaliar_resultado_cadastro('Nome: ANA MARIA\n123456789012345','ANA MARIA','11144477735')
        self.assertIsNone(c)

if __name__=='__main__':unittest.main()
