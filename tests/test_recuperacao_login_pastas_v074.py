"""Regressão de autenticação e saídas; somente dados sintéticos, sem internet."""
import os
import sys
import queue
import tempfile
import threading
import unittest
import shutil
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'src'))
from app_info import APP_VERSION
from diretorios import diretorio_aplicacao, pastas_de_saida
from planilha import SessaoPlanilha
from sisreg_login import Aplicativo

FIXTURE = ROOT/'tests'/'fixtures'/'hospital_sintetico.xlsx'

class TestesPastasEPrivacidade(unittest.TestCase):
    def test_versao_central(self):
        self.assertEqual(APP_VERSION, '0.7.6')

    def test_exe_congelado_define_raiz_sem_meipass(self):
        with tempfile.TemporaryDirectory() as diretorio:
            fake = Path(diretorio)/'RPA_SISREG.exe'
            with patch.object(sys, 'frozen', True, create=True), patch.object(sys, 'executable', str(fake)):
                self.assertEqual(diretorio_aplicacao(), fake.resolve().parent)

    def test_análises_e_logs_criadas_ao_lado_exe(self):
        with tempfile.TemporaryDirectory() as d:
            raiz=Path(d)
            analises, logs = pastas_de_saida(raiz)
            self.assertEqual(analises, raiz/'Análises')
            self.assertEqual(logs, raiz/'Logs')
            self.assertTrue(analises.is_dir() and logs.is_dir())

    def test_xlsx_criado_em_analises_fonte_intacta_e_log_direcionado(self):
        with tempfile.TemporaryDirectory() as d:
            raiz=Path(d)
            origem=raiz/'recebidos'/'modelo.xlsx'
            origem.parent.mkdir()
            shutil.copy2(FIXTURE,origem)
            analises, logs = pastas_de_saida(raiz/'app')
            anterior=origem.read_bytes()
            s=SessaoPlanilha()
            s.carregar(origem)
            copia=s.preparar_copia_de_trabalho(diretorio_saida=analises)
            self.assertEqual(copia.parent, analises)
            s.salvar_na_planilha('teste')
            self.assertEqual(origem.read_bytes(), anterior)
            self.assertTrue(copia.exists())
            self.assertTrue(logs.is_dir())
            self.assertNotEqual(copia.parent,origem.parent)

    def test_falha_nao_exclui_saida_gravada(self):
        with tempfile.TemporaryDirectory() as d:
            raiz=Path(d)
            origem=raiz/'modelo.xlsx';shutil.copy2(FIXTURE,origem)
            analises,logs=pastas_de_saida(raiz/'app')
            s=SessaoPlanilha();s.carregar(origem)
            c=s.preparar_copia_de_trabalho(analises)
            self.assertTrue(s.descartar_copia_nao_gravada())
            self.assertFalse(c.exists())
            s=SessaoPlanilha();s.carregar(origem)
            c=s.preparar_copia_de_trabalho(analises)
            s.salvar_na_planilha('teste')
            self.assertFalse(s.descartar_copia_nao_gravada())
            self.assertTrue(c.exists())

    def test_workflow_compila_exe_renomeado(self):
        wf=(ROOT/'.github'/'workflows'/'compilar-windows.yml').read_text(encoding='utf-8')
        self.assertIn('dist/RPA_SISREG.exe',wf)
        self.assertNotIn('SISREG_Login.exe',wf)
        self.assertIn('RPA_SISREG_Windows',wf)

    def test_creditos_e_licenca(self):
        self.assertIn('bruno.vinicius@joinville.sc.gov.br',(ROOT/'README.md').read_text(encoding='utf-8'))
        self.assertIn('MIT License',(ROOT/'LICENSE').read_text(encoding='utf-8'))
        self.assertIn('LGPD',(ROOT/'SECURITY.md').read_text(encoding='utf-8'))

class FakeBrowser:
    def __init__(self):
        self.context=None
        self.closed=False
    def new_context(self,**kw):
        self.context=FakeContext()
        return self.context
    def close(self):self.closed=True
    def is_connected(self):return not self.closed

class FakeContext:
    def __init__(self):self.pages=[]
    def new_page(self):
        p=FakePage();self.pages.append(p);return p

class FakePage:
    def __init__(self):self.url='https://sisregiii.saude.gov.br/'
    def goto(self,*a,**kw):pass
    def is_closed(self):return False

class FakePlaywright:
    def __init__(self):self.launched=[]
    def __enter__(self):return SimpleNamespace(chromium=SimpleNamespace(launch=self.launch))
    def __exit__(self,*args):return False
    def launch(self,**kw):
        b=FakeBrowser();self.launched.append(b);return b

class TestesLogin(unittest.TestCase):
    def objeto_sem_tk(self):
        app=Aplicativo.__new__(Aplicativo)
        app.parar=threading.Event()
        app.comandos=queue.Queue()
        app.eventos=queue.Queue()
        app.reiniciar_navegador=threading.Event()
        return app

    def test_login_errado_encerra_thread_e_contexto_automaticamente(self):
        pl=FakePlaywright();app=self.objeto_sem_tk()
        with patch('sisreg_login.sync_playwright',return_value=pl),\
             patch('sisreg_login.submeter_login') as login,\
             patch('sisreg_login.desafio_visivel',return_value=True),\
             patch.object(app,'_esperar_intervencao',return_value=False),\
             patch('sisreg_login.autenticao_ativa',return_value=False):
            app._navegador_worker('user_demo','senha_incorreta','chrome')
        self.assertEqual(len(pl.launched),1)
        self.assertTrue(pl.launched[0].closed)
        resultados=[]
        while not app.eventos.empty(): resultados.append(app.eventos.get_nowait())
        self.assertIn('login_enviado',[e[0] for e in resultados])
        self.assertIn('navegador_parou',[e[0] for e in resultados])
        self.assertNotIn('senha_incorreta',repr(resultados))

    def test_login_pode_ser_reexecutado_na_mesma_instancia_sem_preparo(self):
        pl=FakePlaywright();app=self.objeto_sem_tk()
        with patch('sisreg_login.sync_playwright',return_value=pl),\
             patch('sisreg_login.submeter_login'),\
             patch('sisreg_login.desafio_visivel',return_value=True),\
             patch.object(app,'_esperar_intervencao',return_value=False),\
             patch('sisreg_login.autenticao_ativa',return_value=False):
            app._navegador_worker('usuario','senha1','chrome')
            app._navegador_worker('usuario','senha2','chrome')
        self.assertEqual(len(pl.launched),2)
        self.assertTrue(all(b.closed for b in pl.launched))

    def test_browser_fechado_na_sessao_valida_sai_da_thread(self):
        pl=FakePlaywright();app=self.objeto_sem_tk()
        # Browser encerrado manualmente durante a espera por comandos.
        def encerrar(browser):
            browser.close()
        original=pl.launch
        def abre(**kwargs):
            b=original(**kwargs)
            app.comandos.put(('nao_comando',None))  # ignora
            threading.Timer(.10,b.close).start()
            return b
        pl.launch=abre
        with patch('sisreg_login.sync_playwright',return_value=pl),\
             patch('sisreg_login.submeter_login'),\
             patch('sisreg_login.desafio_visivel',return_value=False),\
             patch('sisreg_login.autenticao_ativa',return_value=True):
            app._navegador_worker('usuario','senha_valida','chrome')
        self.assertTrue(pl.launched[0].closed)
        eventos=[]
        while not app.eventos.empty():eventos.append(app.eventos.get_nowait()[0])
        self.assertIn('navegador_parou',eventos)

if __name__=='__main__':unittest.main()

class TesteMigracaoV074(unittest.TestCase):
    def test_remove_coluna_k_e_migra_numero_solicitacao_para_g(self):
        from openpyxl import load_workbook
        from saida_xlsx import adequar_colunas, CAMPOS_FINAIS, COLUNAS_V072
        with tempfile.TemporaryDirectory() as tmp:
            fonte=Path(tmp)/'modelo.xlsx'
            shutil.copy2(FIXTURE,fonte)
            wb=load_workbook(fonte)
            for sh in wb:
                for n,h in enumerate(COLUNAS_V072,8):sh.cell(1,n,h)
                sh.cell(2,7).value=None
                sh.cell(2,11).value='123456789'
                adequar_colunas(sh)
                self.assertEqual([sh.cell(1,c).value for c in range(8,11)], list(CAMPOS_FINAIS))
                self.assertIsNone(sh.cell(1,11).value)
                self.assertEqual(sh.cell(2,7).value,'123456789')
                self.assertEqual(sh.auto_filter.ref, f'A1:J{max(2,sh.max_row)}')
            wb.close()

    def test_nao_sobrescreve_g_preexistente_ao_eliminar_k(self):
        from openpyxl import load_workbook
        from saida_xlsx import adequar_colunas, COLUNAS_V072
        wb=load_workbook(FIXTURE)
        sh=wb['POA CIRURGICO']
        for n,h in enumerate(COLUNAS_V072,8):sh.cell(1,n,h)
        sh.cell(2,7).value='333333333'
        sh.cell(2,11).value='222222222'
        adequar_colunas(sh)
        self.assertEqual(sh.cell(2,7).value,'333333333')
        self.assertIsNone(sh.cell(2,11).value)
        wb.close()

    def test_k_inesperado_bloqueia_descarte_para_proteger_dados(self):
        from openpyxl import load_workbook
        from saida_xlsx import adequar_colunas, COLUNAS_V072, LayoutPlanilhaIncompativel
        wb=load_workbook(FIXTURE)
        sh=wb['POA CIRURGICO']
        for n,h in enumerate(COLUNAS_V072,8):sh.cell(1,n,h)
        sh.cell(2,11).value='VALOR NAO NUMERICO'
        with self.assertRaises(LayoutPlanilhaIncompativel):adequar_colunas(sh)
        wb.close()
