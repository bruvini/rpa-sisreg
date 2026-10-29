import sys
import tempfile
import unittest
import shutil
from pathlib import Path
from types import SimpleNamespace
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'src'))
from sisreg_hospital import (obter_carater, obter_autorizador, obter_procedimento,
    obter_detalhes, ResultadoHospital, selecionar_linha_aih)
from sisreg_navigation import avaliar_cadastro_com_cns_em_negrito
from planilha import SessaoPlanilha

FIXTURES = Path(__file__).resolve().parent/'fixtures'

class TesteDadosHospitalares(unittest.TestCase):
    def test_carater_eletivo(self):
        self.assertEqual(obter_carater('10 - Eletivo'),'10')
    def test_carater_com_rotulo(self):
        self.assertEqual(obter_carater('CARÁTER DA INTERNAÇÃO: 10 – ELETIVO'),'10')
    def test_carater_desconhecido_nao_deduzido(self):
        self.assertEqual(obter_carater('SOLICITAÇÃO DE INTERNAÇÃO'), '')
    def test_carater_01(self):
        self.assertEqual(obter_carater('1 - URGÊNCIA'),'01')
    def test_autorizador_unico(self):
        self.assertEqual(obter_autorizador('Código: 0004532108'),'0004532108')
    def test_autorizador_ambiguo(self):
        self.assertEqual(obter_autorizador('123456 e 345678'), '')
    def test_procedimento_10_digitos(self):
        self.assertEqual(obter_procedimento('0409060178 - HISTERECTOMIA'), '0409060178')
    def test_procedimento_multiplos_codigos(self):
        self.assertEqual(obter_procedimento('0409060178 e 0409040193'), '')
    def test_cns_resultado_destacado_somente_com_identidade(self):
        n,m=avaliar_cadastro_com_cns_em_negrito('Paciente: ANA MARIA EXEMPLO', ['900000000000001'],
                                                 'ANA MARIA EXEMPLO', '11144477735', '11144477735')
        self.assertEqual(n,'900000000000001')
        self.assertEqual(m,'CNS_CONFIRMADO_TD_NEGRITO')
    def test_cns_falso_quando_nome_nao_corresponde(self):
        n,_=avaliar_cadastro_com_cns_em_negrito('Paciente: OUTRA PESSOA', ['900000000000001'],
                                                 'ANA MARIA EXEMPLO', '11144477735', '11144477735')
        self.assertIsNone(n)
    def test_cns_nao_aceita_cpf_outra_pessoa(self):
        n,_=avaliar_cadastro_com_cns_em_negrito('Paciente: ANA MARIA EXEMPLO', ['900000000000001'],
                                                 'ANA MARIA EXEMPLO', '11144477735', '12345678901')
        self.assertIsNone(n)
    def test_cns_ambiguo(self):
        n,_=avaliar_cadastro_com_cns_em_negrito('Paciente: ANA MARIA EXEMPLO',
              ['900000000000001','700000000000001'], 'ANA MARIA EXEMPLO', '11144477735', '11144477735')
        self.assertIsNone(n)

class TesteGravacaoNaOrigem(unittest.TestCase):
    def setUp(self):
        self.d=tempfile.TemporaryDirectory();self.addCleanup(self.d.cleanup)
        self.p=Path(self.d.name)/'hospital.xlsx'
        shutil.copy2(FIXTURES/'hospital_sintetico.xlsx',self.p)
        from openpyxl import load_workbook
        w=load_workbook(self.p)
        ws=w['PT6'];ws['F2']='';ws['G2']='';w.save(self.p);w.close()
        self.s=SessaoPlanilha();self.s.carregar(self.p)
    def test_carater_autorizador_gravados_e_procedimento_nao_sobrescrito(self):
        r=self.s.registros[0]
        origem=r.procedimento
        res=ResultadoHospital('FICHA_CONFERIDA','AIH única','10','0004532108',origem)
        self.s.registrar_hospital(r,res,1.4)
        escrito=self.s.salvar_na_planilha(motivo='teste')
        self.assertEqual(escrito['hospital'],2)
        from openpyxl import load_workbook
        w=load_workbook(self.p,read_only=True)
        ws=w[r.aba]
        self.assertEqual(str(ws.cell(r.linha,6).value),'10')
        self.assertEqual(ws.cell(r.linha,7).value,'0004532108')
        self.assertEqual(str(ws.cell(r.linha,3).value),origem)
        self.assertIsNone(ws.cell(1,11).value)
        self.assertTrue(self.p.with_name('hospital.backup_antes_RPA.xlsx').exists())
        w.close()
    def test_proc_diferente_revisar_sem_sobrescrever(self):
        r=self.s.registros[0]; original=r.procedimento
        res=ResultadoHospital('PROCEDIMENTO_DIVERGENTE','Revisar','10','0004532108','0409040193')
        self.s.registrar_hospital(r,res,3)
        self.assertEqual(r.situacao,'PENDENTE_PROCEDIMENTO_DIVERGENTE')
        self.assertEqual(r.procedimento,original)
    def test_dois_autorizadores_conflitantes_nao_substituem(self):
        r=self.s.registros[0];r.autorizador='00012345'
        self.s.registrar_hospital(r,ResultadoHospital('FICHA_CONFERIDA','', '10','99999999',r.procedimento),1)
        self.assertIn('PENDENTE_AUTORIZADOR_CONFLITANTE',r.situacao)
        self.assertEqual(r.autorizador,'00012345')
    def test_ficha_desconhecida_nao_preenche(self):
        r=self.s.registros[0]
        self.s.registrar_hospital(r,ResultadoHospital('AIH_NAO_COMPROVADA_FICHA',''),1)
        self.assertFalse(r.carater)
        self.assertFalse(r.autorizador)

if __name__=='__main__':unittest.main()

class _Celulas:
    def __init__(self, valores):self.valores=valores;self.clicks=[]
    def count(self):return len(self.valores)
    def all_inner_texts(self):return self.valores
    def nth(self,i):return _Celula(self,i)
class _Celula:
    def __init__(self,grupo,i):self.grupo=grupo;self.i=i
    def is_visible(self):return True
    def click(self,**kwargs):self.grupo.clicks.append(self.i)
class _Linha:
    def __init__(self,valores):self.tds=_Celulas(valores)
    def locator(self,selector):return self.tds
class _Lista:
    def __init__(self,itens):self.itens=itens
    def count(self):return len(self.itens)
    def nth(self,i):return self.itens[i]
class _Pagina:
    url='https://sisregiii.saude.gov.br/cgi-bin/index'
    def __init__(self,linhas):self.frames=[self];self.linhas=linhas
    def locator(self,selector):
        if selector=='#main_page table tr':return _Lista(self.linhas)
        raise AssertionError(selector)

class TestesSelecaoExataAIH(unittest.TestCase):
    def test_abre_aihexata_nao_primeira_linha(self):
        r1=_Linha(['12.05.2026','ANA','PROCEDIMENTO','APROVADO','000000000000-1','OUTRO'])
        r2=_Linha(['12.05.2026','BRUNA','PROCEDIMENTO','APROVADO','000000000000-2','OUTRO'])
        r3=_Linha(['12.05.2026','CARLA','PROCEDIMENTO','APROVADO','900000000012-3','OUTRO'])
        selecionar_linha_aih(_Pagina([r1,r2,r3]),'9000000000123',espera_ms=100)
        self.assertEqual(r1.tds.clicks,[])
        self.assertEqual(r2.tds.clicks,[])
        self.assertEqual(r3.tds.clicks,[4])

    def test_aih_repetida_bloqueia_abertura(self):
        a=_Linha(['A','B','C','D','900000000012-3'])
        b=_Linha(['A','B','C','D','9000000000123'])
        from sisreg_navigation import NavegacaoInesperada
        with self.assertRaises(NavegacaoInesperada):
            selecionar_linha_aih(_Pagina([a,b]),'9000000000123',espera_ms=100)
        self.assertEqual(a.tds.clicks,[])
        self.assertEqual(b.tds.clicks,[])

class _TextoLoc:
    def __init__(self,txt):self.txt=txt
    def inner_text(self,**kwargs):return self.txt
class _FichePage:
    url='https://sisregiii.saude.gov.br/cgi-bin/index'
    frames=[]

class TesteFichaHospitalar(unittest.TestCase):
    def test_ficha_aih_diferente_recusada(self):
        from unittest.mock import patch
        def loc(page,selector,**kwargs):return page,_TextoLoc('AIH 900000000012-4')
        with patch('sisreg_hospital._encontrar',loc):
            r=obter_detalhes(_FichePage(),'9000000000123','0408050063',espera_ficha_ms=25)
        self.assertEqual(r.situacao,'AIH_NAO_COMPROVADA_FICHA')

    def test_ficha_codigo_procedimento_diferente_sinalizado(self):
        from unittest.mock import patch
        d={
          '#fichaInternacao':'AIH 900000000012-3 NOME TESTE',
          '#fichaInternacao > table > tbody > tr:nth-child(46) > td:nth-child(1)':'10 - ELETIVO',
          '#fichaInternacao > table > tbody > tr:nth-child(1) > td:nth-child(1) > b':'0004532108',
          '#fichaInternacao > table > tbody > tr:nth-child(50) > td:nth-child(2)':'0409040193',
        }
        def loc(page,selector,**kwargs):return page,_TextoLoc(d[selector])
        with patch('sisreg_hospital._encontrar',loc):
            r=obter_detalhes(_FichePage(),'9000000000123','0408050063',espera_ficha_ms=25)
        self.assertEqual(r.situacao,'PROCEDIMENTO_DIVERGENTE')
        self.assertEqual(r.carater,'10')
        self.assertEqual(r.solicitacao,'0004532108')
        self.assertEqual(r.autorizador,'')  # Código da solicitação não é autorizador
